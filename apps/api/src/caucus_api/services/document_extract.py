"""Estrazione testo dai documenti caricati (docx, pdf).

v1 deliberatamente senza OCR: un PDF scansionato (nessun layer testuale)
viene rifiutato con errore chiaro invece di produrre un'analisi su un
documento vuoto — meglio nessuna risposta che una risposta non ancorata.
"""

from __future__ import annotations

import io
from dataclasses import dataclass

# Cap di ingestione: oltre questo limite il testo viene tagliato e il flag
# `truncated` viaggia fino al prompt, dove il modello è tenuto a dichiararlo.
MAX_TEXT_CHARS = 300_000

_DOCX_MIME = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"


class DocumentExtractionError(ValueError):
    """Errore utente (formato non supportato, PDF senza testo, file corrotto)."""


@dataclass(frozen=True, slots=True)
class ExtractedDocument:
    text: str
    truncated: bool
    media_type: str


def extract_text(filename: str, data: bytes) -> ExtractedDocument:
    """Estrae il testo da un file docx o pdf, identificato dall'estensione."""
    name = filename.lower()
    if name.endswith(".docx"):
        text = _extract_docx(data)
        media_type = _DOCX_MIME
    elif name.endswith(".pdf"):
        text = _extract_pdf(data)
        media_type = "application/pdf"
    else:
        raise DocumentExtractionError(
            "Formato non supportato: carica un file .docx o .pdf "
            "(per .doc legacy: risalvarlo come .docx)."
        )

    text = text.strip()
    if not text:
        raise DocumentExtractionError(
            "Il documento non contiene testo estraibile. Se è una scansione, "
            "serve un PDF con layer testuale (OCR non ancora supportato)."
        )
    truncated = len(text) > MAX_TEXT_CHARS
    if truncated:
        text = text[:MAX_TEXT_CHARS].rstrip()
    return ExtractedDocument(text=text, truncated=truncated, media_type=media_type)


def _extract_docx(data: bytes) -> str:
    import zipfile

    from docx import Document
    from docx.opc.exceptions import PackageNotFoundError

    try:
        doc = Document(io.BytesIO(data))
    except (PackageNotFoundError, zipfile.BadZipFile, KeyError, ValueError) as exc:
        raise DocumentExtractionError("File .docx non leggibile o corrotto.") from exc

    parts: list[str] = [p.text for p in doc.paragraphs if p.text.strip()]
    # Le tabelle (comuni nei contratti: canoni, scadenze) diventano righe
    # "cella | cella": perdono la grafica ma non il contenuto.
    for table in doc.tables:
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells]
            if any(cells):
                parts.append(" | ".join(cells))
    return "\n".join(parts)


def _extract_pdf(data: bytes) -> str:
    from pypdf import PdfReader
    from pypdf.errors import PdfReadError

    try:
        reader = PdfReader(io.BytesIO(data))
        if reader.is_encrypted:
            raise DocumentExtractionError("PDF protetto da password: rimuovere la protezione.")
        pages = [page.extract_text() or "" for page in reader.pages]
    except DocumentExtractionError:
        raise
    except (PdfReadError, ValueError, KeyError) as exc:
        raise DocumentExtractionError("File .pdf non leggibile o corrotto.") from exc
    return "\n\n".join(p.strip() for p in pages if p.strip())
