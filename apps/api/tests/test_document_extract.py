"""Test estrazione testo documenti (docx, pdf) per l'analisi documentale."""

from __future__ import annotations

import io

import pytest
from docx import Document

from caucus_api.services.document_extract import (
    MAX_TEXT_CHARS,
    DocumentExtractionError,
    extract_text,
)


def _docx_bytes(paragraphs: list[str], table: list[list[str]] | None = None) -> bytes:
    doc = Document()
    for p in paragraphs:
        doc.add_paragraph(p)
    if table:
        t = doc.add_table(rows=len(table), cols=len(table[0]))
        for i, row in enumerate(table):
            for j, cell in enumerate(row):
                t.cell(i, j).text = cell
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def _minimal_pdf(text: str) -> bytes:
    """PDF minimale valido con una riga di testo, xref calcolato."""
    content = f"BT /F1 12 Tf 72 720 Td ({text}) Tj ET".encode()
    objects = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        (
            b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] "
            b"/Contents 4 0 R /Resources << /Font << /F1 5 0 R >> >> >>"
        ),
        b"<< /Length " + str(len(content)).encode() + b" >>\nstream\n" + content + b"\nendstream",
        b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
    ]
    out = io.BytesIO()
    out.write(b"%PDF-1.4\n")
    offsets = []
    for i, obj in enumerate(objects, start=1):
        offsets.append(out.tell())
        out.write(f"{i} 0 obj\n".encode() + obj + b"\nendobj\n")
    xref_pos = out.tell()
    out.write(f"xref\n0 {len(objects) + 1}\n".encode())
    out.write(b"0000000000 65535 f \n")
    for off in offsets:
        out.write(f"{off:010d} 00000 n \n".encode())
    out.write(
        f"trailer\n<< /Size {len(objects) + 1} /Root 1 0 R >>\nstartxref\n{xref_pos}\n%%EOF".encode()
    )
    return out.getvalue()


def test_docx_paragraphs_and_tables():
    data = _docx_bytes(
        ["Contratto di locazione", "Art. 1 — Oggetto"],
        table=[["Canone", "800 €"], ["Scadenza", "31/12/2026"]],
    )
    doc = extract_text("contratto.docx", data)
    assert "Contratto di locazione" in doc.text
    assert "Canone | 800 €" in doc.text
    assert not doc.truncated


def test_pdf_with_text_layer():
    doc = extract_text("atto.pdf", _minimal_pdf("Clausola 5.2: canone mensile"))
    assert "Clausola 5.2" in doc.text
    assert doc.media_type == "application/pdf"


def test_empty_docx_is_rejected():
    with pytest.raises(DocumentExtractionError, match="testo estraibile"):
        extract_text("vuoto.docx", _docx_bytes([]))


def test_unsupported_extension_is_rejected():
    with pytest.raises(DocumentExtractionError, match="Formato non supportato"):
        extract_text("contratto.odt", b"x")


def test_corrupt_docx_is_rejected():
    with pytest.raises(DocumentExtractionError, match="non leggibile"):
        extract_text("rotto.docx", b"not a zip archive")


def test_truncation_flag_at_cap():
    data = _docx_bytes(["x" * 10_000] * 40)  # ~400k > cap
    doc = extract_text("lungo.docx", data)
    assert doc.truncated
    assert len(doc.text) <= MAX_TEXT_CHARS
